#-*- conding: utf-8 -*-
'''
解析pdf里的table
'''
def parse_pdf_table():
    import tabula

    df = tabula.read_pdf(r'C:\Users\20945\Desktop\Spider\local_spider\Other\1.pdf', encoding='gbk', pages='all')
    print(df)

    for j in df:
        for indexs in j.index:
            print(j.loc[indexs].values[1].strip())
'''
解析docx的数据
'''
def parse_docx():
    import docx
    data = docx.Document(r'C:\Users\20945\Desktop\Spider\now_spider\goverment_infomation_source_data\files\杭州市2016年第一批浙江省科技型企业公示名单.docx')
    print(data)
    all_tables = data.tables
    sum = []
    for ta in all_tables:
        rows = ta.rows
        for j in range(1, len(rows)):
            name = ta.cell(j, 1).text
            area = ta.cell(j, 2).text
            year = '2016'
            type = '省科技型'
            batch = '第一批'
            print((name,area))
            sum.append((name,area,year,type,batch))
'''
解析excle里的数据
'''
def parse_excle():
    from xlrd import open_workbook
    workbook = open_workbook(
        r'C:\Users\20945\Desktop\Spider\now_spider\goverment_infomation_source_data\files\附件2-高新区2010-2018年认定的市级科技型初创企业培育工程企业（雏鹰、青蓝）名单 (1).xls')  # 打开excel文件
    sheet2 = workbook.sheet_by_index(0)
    print(sheet2.nrows)
    sum = []
    for i in range(1, sheet2.nrows):
        name = sheet2.cell(i, 1).value
        area = sheet2.cell(i, 2).value
        year = sheet2.cell(i, 3).value
        type = sheet2.cell(i, 4).value
        print((name, area, int(year), type))
        sum.append((name, area, int(year), type))